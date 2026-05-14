"""Generate downloadable documents from agent-provided content."""

from __future__ import annotations

import csv
import json
import re
import textwrap
import zipfile
from dataclasses import asdict, dataclass
from datetime import datetime, timezone
from io import StringIO
from pathlib import Path
from uuid import uuid4

from src.config import settings


SUPPORTED_DOCUMENT_TYPES = {"markdown", "word", "excel", "pdf"}
EXTENSIONS = {"markdown": ".md", "word": ".docx", "excel": ".xlsx", "pdf": ".pdf"}
MIME_TYPES = {
    "markdown": "text/markdown; charset=utf-8",
    "word": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    "excel": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    "pdf": "application/pdf",
}


@dataclass(frozen=True, slots=True)
class DocumentAttachment:
    file_id: str
    filename: str
    document_type: str
    mime_type: str
    download_url: str
    size: int
    created_at: str

    def to_dict(self) -> dict[str, str | int]:
        return asdict(self)


class DocumentGeneratorService:
    """Create local downloadable documents.

    The service intentionally accepts a plain string so the LLM can decide how
    to shape content before tool invocation, while file generation remains
    deterministic and testable.
    """

    def __init__(self, output_directory: str | Path | None = None) -> None:
        self.output_directory = Path(output_directory or settings.generated_documents.directory)
        self.output_directory.mkdir(parents=True, exist_ok=True)

    def generate(self, *, content: str, document_type: str, filename: str | None = None) -> DocumentAttachment:
        normalized_content = content.strip()
        if not normalized_content:
            raise ValueError("content 不能为空")

        normalized_type = document_type.strip().lower()
        if normalized_type not in SUPPORTED_DOCUMENT_TYPES:
            supported = ", ".join(sorted(SUPPORTED_DOCUMENT_TYPES))
            raise ValueError(f"document_type 仅支持: {supported}")

        file_id = uuid4().hex
        final_filename = self._normalize_filename(filename, document_type=normalized_type, file_id=file_id)
        path = self.output_directory / f"{file_id}_{final_filename}"

        if normalized_type == "markdown":
            self._write_markdown(path, normalized_content)
        elif normalized_type == "word":
            self._write_word(path, normalized_content)
        elif normalized_type == "excel":
            self._write_excel(path, normalized_content)
        else:
            self._write_pdf(path, normalized_content)

        return DocumentAttachment(
            file_id=file_id,
            filename=final_filename,
            document_type=normalized_type,
            mime_type=MIME_TYPES[normalized_type],
            download_url=f"/generated-documents/{file_id}/download",
            size=path.stat().st_size,
            created_at=datetime.now(timezone.utc).isoformat(),
        )

    def get_file_path(self, file_id: str) -> Path | None:
        if not re.fullmatch(r"[a-f0-9]{32}", file_id):
            return None
        matches = sorted(self.output_directory.glob(f"{file_id}_*"))
        return matches[0] if matches else None

    @staticmethod
    def _normalize_filename(filename: str | None, *, document_type: str, file_id: str) -> str:
        extension = EXTENSIONS[document_type]
        raw_name = (filename or "").strip()
        if not raw_name:
            raw_name = f"generated-{document_type}-{file_id[:8]}"
        raw_name = Path(raw_name).name
        raw_name = re.sub(r"[\\/:*?\"<>|]+", "-", raw_name)
        raw_name = re.sub(r"\s+", " ", raw_name).strip(" .")
        if not raw_name:
            raw_name = f"generated-{document_type}-{file_id[:8]}"

        path = Path(raw_name)
        if path.suffix.lower() != extension:
            raw_name = f"{path.stem or raw_name}{extension}"
        return raw_name[:180]

    @staticmethod
    def _write_markdown(path: Path, content: str) -> None:
        path.write_text(content + "\n", encoding="utf-8")

    @staticmethod
    def _write_word(path: Path, content: str) -> None:
        try:
            from docx import Document
        except ImportError as exc:  # pragma: no cover - dependency guard
            raise RuntimeError("生成 Word 文档需要安装 python-docx") from exc

        document = Document()
        rows_buffer: list[list[str]] = []

        def flush_table() -> None:
            nonlocal rows_buffer
            rows = _normalize_markdown_table(rows_buffer)
            if not rows:
                rows_buffer = []
                return
            table = document.add_table(rows=len(rows), cols=max(len(row) for row in rows))
            table.style = "Table Grid"
            for row_index, row in enumerate(rows):
                for col_index, cell in enumerate(row):
                    table.cell(row_index, col_index).text = cell
            rows_buffer = []

        for line in content.splitlines():
            stripped = line.strip()
            if _looks_like_markdown_table_row(stripped):
                rows_buffer.append(_split_markdown_table_row(stripped))
                continue
            flush_table()
            if not stripped:
                continue
            heading = re.match(r"^(#{1,6})\s+(.+)$", stripped)
            if heading:
                document.add_heading(heading.group(2).strip(), level=min(len(heading.group(1)), 4))
            elif re.match(r"^[-*+]\s+", stripped):
                document.add_paragraph(re.sub(r"^[-*+]\s+", "", stripped), style="List Bullet")
            elif re.match(r"^\d+[.)]\s+", stripped):
                document.add_paragraph(re.sub(r"^\d+[.)]\s+", "", stripped), style="List Number")
            else:
                document.add_paragraph(stripped)
        flush_table()
        document.save(path)

    @staticmethod
    def _write_excel(path: Path, content: str) -> None:
        try:
            from openpyxl import Workbook
        except ImportError as exc:  # pragma: no cover - dependency guard
            raise RuntimeError("生成 Excel 文档需要安装 openpyxl") from exc

        workbook = Workbook()
        worksheet = workbook.active
        worksheet.title = "Sheet1"
        rows = _parse_tabular_content(content)
        if not rows:
            rows = [["内容"], *[[line] for line in content.splitlines() if line.strip()]]
        for row in rows:
            worksheet.append(row)
        for column_cells in worksheet.columns:
            max_length = max((len(str(cell.value or "")) for cell in column_cells), default=8)
            worksheet.column_dimensions[column_cells[0].column_letter].width = min(max(max_length + 2, 10), 60)
        workbook.save(path)

    @staticmethod
    def _write_pdf(path: Path, content: str) -> None:
        title, body = _extract_title_and_body(content)
        lines: list[str] = []
        if title:
            lines.append(title)
            lines.append("")
        for raw_line in body.splitlines():
            normalized = re.sub(r"^#{1,6}\s+", "", raw_line).strip()
            if not normalized:
                lines.append("")
                continue
            wrapped = textwrap.wrap(normalized, width=74) or [normalized]
            lines.extend(wrapped)
        _write_basic_pdf(path, lines)


def _looks_like_markdown_table_row(line: str) -> bool:
    return line.startswith("|") and line.endswith("|") and line.count("|") >= 2


def _split_markdown_table_row(line: str) -> list[str]:
    return [cell.strip() for cell in line.strip("|").split("|")]


def _normalize_markdown_table(rows: list[list[str]]) -> list[list[str]]:
    normalized = []
    for row in rows:
        if row and all(re.fullmatch(r":?-{3,}:?", cell.strip()) for cell in row):
            continue
        normalized.append(row)
    return normalized


def _parse_tabular_content(content: str) -> list[list[str]]:
    stripped = content.strip()
    if not stripped:
        return []

    try:
        parsed = json.loads(stripped)
    except json.JSONDecodeError:
        parsed = None
    if isinstance(parsed, list) and parsed:
        if all(isinstance(item, dict) for item in parsed):
            headers = list(dict.fromkeys(key for item in parsed for key in item.keys()))
            return [headers, *[[str(item.get(header, "")) for header in headers] for item in parsed]]
        if all(isinstance(item, list) for item in parsed):
            return [[str(cell) for cell in item] for item in parsed]

    table_rows = [_split_markdown_table_row(line.strip()) for line in stripped.splitlines() if _looks_like_markdown_table_row(line.strip())]
    table_rows = _normalize_markdown_table(table_rows)
    if table_rows:
        return table_rows

    try:
        csv_rows = list(csv.reader(StringIO(stripped)))
    except csv.Error:
        csv_rows = []
    if len(csv_rows) > 1 or (csv_rows and len(csv_rows[0]) > 1):
        return [[cell.strip() for cell in row] for row in csv_rows]

    return []


def _extract_title_and_body(content: str) -> tuple[str, str]:
    lines = content.splitlines()
    for index, line in enumerate(lines):
        heading = re.match(r"^#{1,6}\s+(.+)$", line.strip())
        if heading:
            return heading.group(1).strip(), "\n".join(lines[:index] + lines[index + 1 :])
    return "", content


def _pdf_hex_text(value: str) -> str:
    return value.encode("utf-16-be", errors="replace").hex().upper()


def _write_basic_pdf(path: Path, lines: list[str]) -> None:
    visible_lines = lines or ["Generated document"]
    page_height = 842
    line_height = 15
    top = 800
    bottom = 54
    pages = [visible_lines[index : index + 48] for index in range(0, len(visible_lines), 48)] or [[]]
    objects: list[bytes] = []

    def add_object(payload: str | bytes) -> int:
        data = payload.encode("latin-1", errors="replace") if isinstance(payload, str) else payload
        objects.append(data)
        return len(objects)

    page_ids: list[int] = []
    content_ids: list[int] = []
    for page_lines in pages:
        y = top
        commands = ["BT", "/F1 11 Tf", "50 800 Td"]
        first = True
        for line in page_lines:
            if not first:
                commands.append(f"0 -{line_height} Td")
                y -= line_height
            first = False
            if y < bottom:
                break
            commands.append(f"<{_pdf_hex_text(line)}> Tj")
        commands.append("ET")
        stream = "\n".join(commands).encode("latin-1", errors="replace")
        content_id = add_object(f"<< /Length {len(stream)} >>\nstream\n".encode("latin-1") + stream + b"\nendstream")
        content_ids.append(content_id)
        page_id = add_object(b"")
        page_ids.append(page_id)

    pages_id = len(objects) + 1
    font_id = pages_id + 1
    cid_font_id = pages_id + 2
    descriptor_id = pages_id + 3
    catalog_id = pages_id + 4

    for index, page_id in enumerate(page_ids):
        objects[page_id - 1] = (
            f"<< /Type /Page /Parent {pages_id} 0 R /MediaBox [0 0 595 {page_height}] "
            f"/Resources << /Font << /F1 {font_id} 0 R >> >> /Contents {content_ids[index]} 0 R >>"
        ).encode("latin-1")

    kids = " ".join(f"{page_id} 0 R" for page_id in page_ids)
    add_object(f"<< /Type /Pages /Kids [{kids}] /Count {len(page_ids)} >>")
    add_object(f"<< /Type /Font /Subtype /Type0 /BaseFont /STSong-Light /Encoding /UniGB-UCS2-H /DescendantFonts [{cid_font_id} 0 R] >>")
    add_object(
        f"<< /Type /Font /Subtype /CIDFontType0 /BaseFont /STSong-Light /CIDSystemInfo "
        f"<< /Registry (Adobe) /Ordering (GB1) /Supplement 2 >> /FontDescriptor {descriptor_id} 0 R >>"
    )
    add_object(
        "<< /Type /FontDescriptor /FontName /STSong-Light /Flags 4 /FontBBox [0 -120 1000 880] "
        "/ItalicAngle 0 /Ascent 880 /Descent -120 /CapHeight 880 /StemV 80 >>"
    )
    add_object(f"<< /Type /Catalog /Pages {pages_id} 0 R >>")

    output = bytearray(b"%PDF-1.4\n%\xe2\xe3\xcf\xd3\n")
    offsets = [0]
    for object_id, payload in enumerate(objects, start=1):
        offsets.append(len(output))
        output.extend(f"{object_id} 0 obj\n".encode("latin-1"))
        output.extend(payload)
        output.extend(b"\nendobj\n")
    xref_offset = len(output)
    output.extend(f"xref\n0 {len(objects) + 1}\n".encode("latin-1"))
    output.extend(b"0000000000 65535 f \n")
    for offset in offsets[1:]:
        output.extend(f"{offset:010d} 00000 n \n".encode("latin-1"))
    output.extend(
        (
            f"trailer\n<< /Size {len(objects) + 1} /Root {catalog_id} 0 R >>\n"
            f"startxref\n{xref_offset}\n%%EOF\n"
        ).encode("latin-1")
    )
    path.write_bytes(output)
