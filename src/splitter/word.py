"""Word-aware section builders for RAG chunking."""

from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path
import re
from typing import Any

from langchain_core.documents import Document

from src.loader import load_documents

from .base import SplitterDependencyError

_DOCX_FILETYPE = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
_DOC_FILETYPE = "application/msword"
_HEADING_STYLE_PATTERN = re.compile(r"^(heading|标题)\s*(\d+)$", re.IGNORECASE)


@dataclass(frozen=True, slots=True)
class WordBlock:
    """Normalized Word content block."""

    text: str
    kind: str = "paragraph"
    level: int | None = None
    metadata: dict[str, Any] | None = None


def build_word_sections(file_path: str | Path, *, loader_kwargs: dict[str, Any] | None = None) -> list[Document]:
    """Build heading-aware section documents from a Word file."""

    path = Path(file_path)
    if path.suffix.lower() == ".docx":
        blocks = _load_docx_blocks(path)
        return build_word_sections_from_blocks(path, blocks, filetype=_DOCX_FILETYPE)

    documents = load_documents(path, **(loader_kwargs or {}))
    return build_word_sections_from_documents(path, documents)


def build_word_sections_from_blocks(
    path: Path,
    blocks: list[WordBlock],
    *,
    filetype: str,
) -> list[Document]:
    """Group normalized Word blocks into heading-aware section documents."""

    sections: list[Document] = []
    heading_stack: dict[int, str] = {}
    current_lines: list[str] = []
    current_metadata: dict[str, Any] = {}
    current_has_body = False

    def reset_current() -> None:
        nonlocal current_lines, current_metadata, current_has_body
        current_lines = _heading_lines(heading_stack)
        current_metadata = _heading_metadata(path, filetype, heading_stack)
        current_has_body = False

    def flush() -> None:
        nonlocal current_lines, current_metadata, current_has_body
        text = "\n\n".join(line for line in current_lines if line.strip()).strip()
        if not text:
            return
        sections.append(Document(page_content=text, metadata=dict(current_metadata)))
        current_lines = []
        current_metadata = {}
        current_has_body = False

    for block in blocks:
        text = _clean_text(block.text)
        if not text:
            continue

        if block.kind == "heading":
            if current_lines and current_has_body:
                flush()
            level = max(1, min(block.level or 1, 6))
            for existing_level in list(heading_stack):
                if existing_level >= level:
                    heading_stack.pop(existing_level)
            heading_stack[level] = text
            reset_current()
            continue

        if not current_lines:
            reset_current()
        current_lines.append(text)
        current_has_body = True

    if current_lines:
        flush()

    return sections


def build_word_sections_from_documents(path: Path, documents: list[Document]) -> list[Document]:
    """Build sections from LangChain Unstructured Word elements."""

    blocks: list[WordBlock] = []
    for document in documents:
        metadata = dict(document.metadata)
        text = _clean_text(document.page_content)
        if not text:
            continue
        category = str(metadata.get("category") or metadata.get("element_category") or "")
        if category.lower() in {"title", "header"}:
            level = _category_depth(metadata)
            blocks.append(WordBlock(text=text, kind="heading", level=level, metadata=metadata))
        else:
            blocks.append(WordBlock(text=text, kind="paragraph", metadata=metadata))

    return build_word_sections_from_blocks(path, blocks, filetype=_DOC_FILETYPE)


def ensure_word_heading_context(document: Document) -> Document:
    """Prefix heading context to split chunks that lost it."""

    heading_context = str(document.metadata.get("heading_context") or "").strip()
    if not heading_context:
        return document

    content = document.page_content.strip()
    if content.startswith(heading_context):
        return document

    return Document(
        page_content=f"{heading_context}\n\n{content}",
        metadata=dict(document.metadata),
    )


def _load_docx_blocks(path: Path) -> list[WordBlock]:
    try:
        from docx import Document as DocxDocument
        from docx.oxml.table import CT_Tbl
        from docx.oxml.text.paragraph import CT_P
        from docx.table import Table
        from docx.text.paragraph import Paragraph
    except ImportError as exc:
        raise SplitterDependencyError("无法导入 python-docx。请安装依赖: pip install python-docx") from exc

    document = DocxDocument(path)
    blocks: list[WordBlock] = []
    for child in document.element.body.iterchildren():
        if isinstance(child, CT_P):
            paragraph = Paragraph(child, document)
            text = _clean_text(paragraph.text)
            if not text:
                continue
            level = _heading_level(paragraph)
            blocks.append(WordBlock(text=text, kind="heading" if level else "paragraph", level=level))
        elif isinstance(child, CT_Tbl):
            table = Table(child, document)
            text = _table_to_text(table)
            if text:
                blocks.append(WordBlock(text=text, kind="table"))
    return blocks


def _table_to_text(table: Any) -> str:
    rows = []
    for row_index, row in enumerate(table.rows, start=1):
        cells = [_clean_text(cell.text) for cell in row.cells]
        if any(cells):
            rows.append(f"表格第{row_index}行: " + " | ".join(cell or "-" for cell in cells))
    return "\n".join(rows)


def _heading_level(paragraph: Any) -> int | None:
    style = paragraph.style
    candidates = [
        getattr(style, "name", ""),
        getattr(style, "style_id", ""),
    ]
    for candidate in candidates:
        normalized = str(candidate).replace("_", " ").strip()
        match = _HEADING_STYLE_PATTERN.match(normalized)
        if match:
            return int(match.group(2))
    return None


def _category_depth(metadata: dict[str, Any]) -> int:
    for key in ("category_depth", "header_depth", "depth"):
        value = metadata.get(key)
        if value is None:
            continue
        try:
            return max(1, min(int(value) + 1, 6))
        except (TypeError, ValueError):
            continue
    return 1


def _heading_lines(heading_stack: dict[int, str]) -> list[str]:
    return [f"{'#' * level} {heading_stack[level]}" for level in sorted(heading_stack)]


def _heading_metadata(path: Path, filetype: str, heading_stack: dict[int, str]) -> dict[str, Any]:
    metadata: dict[str, Any] = {
        "source": str(path),
        "filename": path.name,
        "filetype": filetype,
        "chunk_type": "word_section",
    }
    heading_lines = _heading_lines(heading_stack)
    if heading_lines:
        metadata["heading_context"] = "\n".join(heading_lines)
    for level, heading in heading_stack.items():
        metadata[f"h{level}"] = heading
    return metadata


def _clean_text(text: str) -> str:
    return "\n".join(line.strip() for line in str(text).splitlines() if line.strip()).strip()
