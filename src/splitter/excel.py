"""Spreadsheet-aware splitting helpers for Excel files."""

from __future__ import annotations

from collections.abc import Iterable
from dataclasses import dataclass
from datetime import date, datetime, time
from pathlib import Path
from typing import Any

from langchain_core.documents import Document

from .base import SplitterDependencyError

_EXCEL_FILETYPES = {
    ".xls": "application/vnd.ms-excel",
    ".xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
}


@dataclass(frozen=True, slots=True)
class ExcelSheet:
    """Normalized worksheet data."""

    name: str
    index: int
    rows: list[list[Any]]


@dataclass(frozen=True, slots=True)
class _TableRows:
    preamble: list[str]
    header: list[str]
    rows: list[tuple[int, list[str]]]
    header_row_number: int


def split_excel_file(
    file_path: str | Path,
    *,
    chunk_size: int,
    chunk_overlap: int,
) -> list[Document]:
    """Split an Excel workbook into table-aware chunks.

    Each chunk keeps the workbook name, sheet name, table header, and one or
    more complete data rows together. This is better suited to RAG than generic
    character splitting because table cells only make sense with their headers.
    """

    path = Path(file_path)
    sheets = _load_workbook(path)
    return _build_excel_documents(
        path,
        sheets,
        chunk_size=chunk_size,
        chunk_overlap=chunk_overlap,
    )


def _build_excel_documents(
    path: Path,
    sheets: Iterable[ExcelSheet],
    *,
    chunk_size: int,
    chunk_overlap: int,
) -> list[Document]:
    documents: list[Document] = []
    for sheet in sheets:
        table = _normalize_sheet_rows(sheet.rows)
        if table is None:
            continue

        documents.extend(
            _chunk_sheet_table(
                path=path,
                sheet=sheet,
                table=table,
                chunk_size=chunk_size,
                chunk_overlap=chunk_overlap,
            )
        )
    return documents


def _chunk_sheet_table(
    *,
    path: Path,
    sheet: ExcelSheet,
    table: _TableRows,
    chunk_size: int,
    chunk_overlap: int,
) -> list[Document]:
    prefix_lines = _prefix_lines(path, sheet, table)
    prefix = "\n".join(prefix_lines)
    filetype = _EXCEL_FILETYPES.get(path.suffix.lower(), "application/vnd.ms-excel")

    if not table.rows:
        return [
            Document(
                page_content=prefix,
                metadata={
                    "source": str(path),
                    "filename": path.name,
                    "filetype": filetype,
                    "sheet_name": sheet.name,
                    "sheet_index": sheet.index,
                    "header_row": table.header_row_number,
                    "chunk_type": "excel_table",
                    "row_count": 0,
                },
            )
        ]

    documents: list[Document] = []
    active_rows: list[tuple[int, str]] = []
    active_len = len(prefix)

    for row_number, values in table.rows:
        row_line = _format_data_row(row_number, table.header, values)
        projected_len = active_len + len(row_line) + 1
        if active_rows and projected_len > chunk_size:
            documents.append(
                _make_excel_document(
                    path=path,
                    filetype=filetype,
                    sheet=sheet,
                    table=table,
                    prefix=prefix,
                    rows=active_rows,
                )
            )
            active_rows = _overlap_rows(active_rows, chunk_overlap)
            active_len = len(prefix) + sum(len(line) + 1 for _, line in active_rows)

        active_rows.append((row_number, row_line))
        active_len += len(row_line) + 1

    if active_rows:
        documents.append(
            _make_excel_document(
                path=path,
                filetype=filetype,
                sheet=sheet,
                table=table,
                prefix=prefix,
                rows=active_rows,
            )
        )

    return documents


def _make_excel_document(
    *,
    path: Path,
    filetype: str,
    sheet: ExcelSheet,
    table: _TableRows,
    prefix: str,
    rows: list[tuple[int, str]],
) -> Document:
    row_numbers = [row_number for row_number, _ in rows]
    content = "\n".join([prefix, *[line for _, line in rows]])
    return Document(
        page_content=content,
        metadata={
            "source": str(path),
            "filename": path.name,
            "filetype": filetype,
            "sheet_name": sheet.name,
            "sheet_index": sheet.index,
            "header_row": table.header_row_number,
            "start_row": min(row_numbers),
            "end_row": max(row_numbers),
            "row_count": len(rows),
            "chunk_type": "excel_table",
        },
    )


def _prefix_lines(path: Path, sheet: ExcelSheet, table: _TableRows) -> list[str]:
    lines = [
        f"文件: {path.name}",
        f"工作表: {sheet.name}",
    ]
    if table.preamble:
        lines.append(f"表格说明: {'；'.join(table.preamble)}")
    lines.append(f"表头: {' | '.join(table.header)}")
    return lines


def _format_data_row(row_number: int, header: list[str], values: list[str]) -> str:
    pairs = [
        f"{column}={value}"
        for column, value in zip(header, values, strict=False)
        if value
    ]
    if not pairs:
        pairs = ["空行"]
    return f"第{row_number}行: " + "；".join(pairs)


def _overlap_rows(rows: list[tuple[int, str]], chunk_overlap: int) -> list[tuple[int, str]]:
    if chunk_overlap <= 0:
        return []

    selected: list[tuple[int, str]] = []
    total = 0
    for row in reversed(rows):
        row_len = len(row[1]) + 1
        if selected and total + row_len > chunk_overlap:
            break
        selected.append(row)
        total += row_len
    selected.reverse()
    return selected


def _normalize_sheet_rows(rows: list[list[Any]]) -> _TableRows | None:
    normalized = [_trim_row([_format_cell(value) for value in row]) for row in rows]
    normalized = _trim_empty_edges(normalized)
    if not normalized:
        return None

    header_index = _detect_header_index(normalized)
    max_columns = max(len(row) for row in normalized[header_index:])
    header = _normalize_header(_pad_row(normalized[header_index], max_columns))
    data_rows = []
    for offset, row in enumerate(normalized[header_index + 1 :], start=header_index + 2):
        padded = _pad_row(row, max_columns)
        if any(padded):
            data_rows.append((offset, padded))

    preamble = [" | ".join(cell for cell in row if cell) for row in normalized[:header_index]]
    return _TableRows(
        preamble=[line for line in preamble if line],
        header=header,
        rows=data_rows,
        header_row_number=header_index + 1,
    )


def _detect_header_index(rows: list[list[str]]) -> int:
    search_limit = min(len(rows), 12)
    best_index = 0
    best_score = -1

    for index, row in enumerate(rows[:search_limit]):
        current_count = sum(1 for cell in row if cell)
        next_count = _next_non_empty_count(rows, index + 1)
        previous_count = sum(1 for cell in rows[index - 1] if cell) if index > 0 else 0
        score = current_count * 2 + min(next_count, current_count) - previous_count
        if current_count >= 2 and score > best_score:
            best_index = index
            best_score = score

    if best_score >= 0:
        return best_index
    return 0


def _next_non_empty_count(rows: list[list[str]], start: int) -> int:
    for row in rows[start:]:
        count = sum(1 for cell in row if cell)
        if count:
            return count
    return 0


def _normalize_header(row: list[str]) -> list[str]:
    seen: dict[str, int] = {}
    header: list[str] = []
    for index, value in enumerate(row, start=1):
        name = value.strip() or f"列{index}"
        count = seen.get(name, 0) + 1
        seen[name] = count
        if count > 1:
            name = f"{name}_{count}"
        header.append(name)
    return header


def _trim_empty_edges(rows: list[list[str]]) -> list[list[str]]:
    start = 0
    end = len(rows)
    while start < end and not any(rows[start]):
        start += 1
    while end > start and not any(rows[end - 1]):
        end -= 1
    return rows[start:end]


def _trim_row(row: list[str]) -> list[str]:
    end = len(row)
    while end > 0 and not row[end - 1]:
        end -= 1
    return row[:end]


def _pad_row(row: list[str], length: int) -> list[str]:
    if len(row) >= length:
        return row[:length]
    return [*row, *([""] * (length - len(row)))]


def _format_cell(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.isoformat(sep=" ", timespec="seconds")
    if isinstance(value, (date, time)):
        return value.isoformat()
    text = str(value).strip()
    if text.endswith(".0"):
        try:
            as_float = float(text)
        except ValueError:
            return text
        if as_float.is_integer():
            return str(int(as_float))
    return text


def _load_workbook(path: Path) -> list[ExcelSheet]:
    suffix = path.suffix.lower()
    if suffix == ".xlsx":
        return _load_xlsx(path)
    if suffix == ".xls":
        return _load_xls(path)
    raise ValueError(f"暂不支持 Excel 扩展名: {suffix}")


def _load_xlsx(path: Path) -> list[ExcelSheet]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise SplitterDependencyError("无法导入 openpyxl。请安装依赖: pip install openpyxl") from exc

    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheets = []
        for index, worksheet in enumerate(workbook.worksheets):
            rows = [list(row) for row in worksheet.iter_rows(values_only=True)]
            sheets.append(ExcelSheet(name=worksheet.title, index=index, rows=rows))
        return sheets
    finally:
        workbook.close()


def _load_xls(path: Path) -> list[ExcelSheet]:
    try:
        import xlrd
    except ImportError as exc:
        raise SplitterDependencyError("无法导入 xlrd。请安装依赖: pip install xlrd") from exc

    workbook = xlrd.open_workbook(path)
    sheets = []
    for index in range(workbook.nsheets):
        sheet = workbook.sheet_by_index(index)
        rows = []
        for row_index in range(sheet.nrows):
            row = []
            for col_index in range(sheet.ncols):
                cell = sheet.cell(row_index, col_index)
                value = cell.value
                if cell.ctype == xlrd.XL_CELL_DATE:
                    value = xlrd.xldate_as_datetime(value, workbook.datemode)
                row.append(value)
            rows.append(row)
        sheets.append(ExcelSheet(name=sheet.name, index=index, rows=rows))
    return sheets
